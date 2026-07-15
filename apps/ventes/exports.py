"""Exports de l'historique des ventes : Excel (openpyxl) et PDF (xhtml2pdf)."""
from io import BytesIO

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone

COLONNES = [
    "Code Vente",
    "Demandeur",
    "Quantité (Kg)",
    "Prix Unitaire (DH/Kg)",
    "Prix Total (DH)",
    "Date",
    "Utilisateur",
]


def _lignes(queryset):
    """Transforme un queryset de ventes en lignes prêtes pour l'export."""
    for v in queryset:
        yield [
            v.code_vente,
            v.demandeur.libelle,
            float(v.quantite),
            float(v.prix_unitaire),
            float(v.prix_total),
            v.date_vente.strftime("%d/%m/%Y"),
            v.utilisateur.get_full_name() or v.utilisateur.username,
        ]


def export_ventes_excel(queryset) -> HttpResponse:
    """Génère un classeur Excel (.xlsx) de l'historique filtré."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventes"

    entete_fill = PatternFill("solid", fgColor="0D6EFD")
    entete_font = Font(bold=True, color="FFFFFF")
    for col, titre in enumerate(COLONNES, start=1):
        cell = ws.cell(row=1, column=col, value=titre)
        cell.fill = entete_fill
        cell.font = entete_font

    for ligne in _lignes(queryset):
        ws.append(ligne)

    # Largeurs de colonnes lisibles.
    largeurs = [20, 30, 14, 18, 16, 12, 20]
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[chr(64 + i)].width = largeur

    flux = BytesIO()
    wb.save(flux)
    flux.seek(0)

    horodatage = timezone.localdate().strftime("%Y%m%d")
    response = HttpResponse(
        flux.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="ventes_{horodatage}.xlsx"'
    return response


def export_ventes_pdf(queryset) -> HttpResponse:
    """Génère un PDF de l'historique filtré via xhtml2pdf."""
    from xhtml2pdf import pisa

    html = render_to_string(
        "ventes/export_pdf.html",
        {
            "ventes": queryset,
            "genere_le": timezone.localtime().strftime("%d/%m/%Y %H:%M"),
        },
    )

    flux = BytesIO()
    resultat = pisa.CreatePDF(src=html, dest=flux, encoding="utf-8")
    if resultat.err:
        return HttpResponse("Erreur lors de la génération du PDF.", status=500)

    horodatage = timezone.localdate().strftime("%Y%m%d")
    response = HttpResponse(flux.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="ventes_{horodatage}.pdf"'
    return response
